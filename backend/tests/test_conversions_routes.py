"""Routes des conversions : import, relecture, colonnes, exports, envoi vers Atelier Tab, IA, exemples."""

import csv
import io
import json
from pathlib import Path
from typing import Any

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.core.config import Settings
from app.schemas.ai import ChatMessage, ProviderName
from app.services.ai import AIProviderError

ROOT = Path(__file__).resolve().parents[2]
CSV = (ROOT / "samples" / "export-banque.csv").read_bytes()
PDF = (ROOT / "samples" / "releve-mars-2026.pdf").read_bytes()


def upload(client: TestClient, content: bytes = CSV, name: str = "export-banque.csv", **options: Any) -> dict[str, Any]:
    data = {"options": json.dumps(options)} if options else None
    r = client.post("/api/conversions", files={"file": (name, content, "application/octet-stream")}, data=data)
    assert r.status_code == 201, r.text
    return r.json()


def make_tab(settings: Settings, projects: int = 1) -> Path:
    """Dossier d'Atelier Tab minimal : un espace, `projects` projets."""
    root = settings.tab_dir
    (root / "tables").mkdir(parents=True)
    index = {
        "format": "atelier_tab",
        "version": 1,
        "workspaceOrder": ["ws_1"],
        "workspaces": {"ws_1": {"id": "ws_1", "name": "Maison", "projectIds": [f"proj_{i}" for i in range(projects)]}},
        "projects": {
            f"proj_{i}": {"id": f"proj_{i}", "workspaceId": "ws_1", "name": f"Budget {i}", "tableIds": []}
            for i in range(projects)
        },
    }
    (root / "atelier.json").write_text(json.dumps(index), encoding="utf-8")
    return root


def test_upload_and_view(client: TestClient) -> None:
    conv = upload(client)
    assert conv["name"] == "export-banque"
    assert conv["detected"] == "Relevé bancaire · 18 opération(s)"
    assert conv["rows"] == 18 and conv["columns"] == 6
    assert conv["source"]["encoding"] == "cp1252" and conv["source"]["kind"] == "csv"
    assert len(conv["preview"]) == 18 and conv["preview_truncated"] is False
    assert conv["statement"]["total_debit"] == 1423.35
    again = client.get(f"/api/conversions/{conv['id']}").json()
    assert again["table_columns"] == conv["table_columns"]
    listed = client.get("/api/conversions").json()
    assert [c["id"] for c in listed] == [conv["id"]] and "preview" not in listed[0]


def test_upload_refusals(client: TestClient, settings: Settings) -> None:
    r = client.post("/api/conversions", files={"file": ("video.mp4", b"\x00\x00\x00 ftypisom", "video/mp4")})
    assert r.status_code == 415 and "Formats acceptés" in r.json()["detail"]
    r = client.post("/api/conversions", files={"file": ("vide.csv", b"", "text/csv")})
    assert r.status_code == 422
    r = client.post("/api/conversions", files={"file": ("a.csv", b"x;y\n1;2\n", "text/csv")}, data={"options": "{oops"})
    assert r.status_code == 422


def test_upload_too_large(settings: Settings) -> None:
    from main import create_app

    small = settings.model_copy(update={"max_upload_mb": 1})
    with TestClient(create_app(small)) as c:
        r = c.post("/api/conversions", files={"file": ("gros.csv", b"a;b\n" * 400_000, "text/csv")})
        assert r.status_code == 413


def test_reparse_with_options(client: TestClient) -> None:
    conv = upload(client)
    r = client.put(f"/api/conversions/{conv['id']}/options", json={"profile": "raw"})
    assert r.status_code == 200
    raw = r.json()
    assert [c["label"] for c in raw["table_columns"]][:3] == ["Date", "Date de valeur", "Libellé"]
    assert raw["statement"] is None and raw["options"]["profile"] == "raw"


def test_update_columns(client: TestClient) -> None:
    conv = upload(client)
    cols = conv["table_columns"]
    cols[2]["label"] = "Opération"
    cols[1]["include"] = False
    cols.insert(0, cols.pop(5))  # Montant en premier
    r = client.put(f"/api/conversions/{conv['id']}/columns", json={"columns": cols})
    assert r.status_code == 200, r.text
    body = r.json()
    assert [c["key"] for c in body["table_columns"]][0] == "montant"
    assert body["columns"] == 5
    # Changement de type : les valeurs sont reconverties.
    cols[0]["type"] = "text"
    body = client.put(f"/api/conversions/{conv['id']}/columns", json={"columns": cols}).json()
    assert body["preview"][0]["montant"] == "-84.2"
    # Clés différentes ou noms en double : refusé.
    bad = [dict(c) for c in cols]
    bad[0]["key"] = "autre"
    assert client.put(f"/api/conversions/{conv['id']}/columns", json={"columns": bad}).status_code == 409
    dup = [dict(c) for c in cols]
    dup[3]["label"] = dup[2]["label"]
    dup[3]["include"] = dup[2]["include"] = True
    assert client.put(f"/api/conversions/{conv['id']}/columns", json={"columns": dup}).status_code == 422


def test_rename_and_delete(client: TestClient, settings: Settings) -> None:
    conv = upload(client)
    assert (
        client.put(f"/api/conversions/{conv['id']}/name", json={"name": "  Relevé   mars "}).json()["name"]
        == "Relevé mars"
    )
    assert client.delete(f"/api/conversions/{conv['id']}").json() == {"deleted": True}
    assert client.get(f"/api/conversions/{conv['id']}").status_code == 404
    assert not list((settings.data_dir / "sources").glob("*"))
    assert client.get("/api/conversions/../../etc").status_code == 404


def test_exports(client: TestClient) -> None:
    conv = upload(client)
    cid = conv["id"]
    r = client.get(f"/api/conversions/{cid}/export", params={"format": "csv-fr"})
    assert r.status_code == 200 and r.content.startswith("﻿".encode())
    assert "filename*=UTF-8''export-banque.csv" in r.headers["content-disposition"]
    rows = list(csv.reader(io.StringIO(r.content.decode("utf-8-sig")), delimiter=";"))
    assert rows[0] == ["Date", "Date de valeur", "Libellé", "Débit", "Crédit", "Montant"]
    assert rows[1] == [
        "02/03/2026",
        "02/03/2026",
        "PRLV SEPA EDF ENERGIE ECH/020326 ID EMETTEUR/FR12ZZZ000001",
        "84,20",
        "",
        "-84,20",
    ]

    en = list(csv.reader(io.StringIO(client.get(f"/api/conversions/{cid}/export?format=csv").text)))
    assert en[1][0] == "2026-03-02" and en[1][3] == "84.20"

    txt = client.get(f"/api/conversions/{cid}/export?format=txt").text
    assert txt.splitlines()[0].split("\t")[0] == "Date"

    book = load_workbook(io.BytesIO(client.get(f"/api/conversions/{cid}/export?format=xlsx").content))
    ws = book.active
    assert ws["A1"].value == "Date" and ws["A2"].value.year == 2026 and ws["D2"].value == 84.2
    assert "€" in ws["D2"].number_format

    table = client.get(f"/api/conversions/{cid}/export?format=json").json()
    assert table["format"] == "atelier-table/1" and table["total"] == 18
    assert {c["key"]: c["type"] for c in table["columns"]}["debit"] == "number"
    assert client.get(f"/api/conversions/{cid}/export?format=docx").status_code == 422


def test_export_skips_excluded_columns(client: TestClient) -> None:
    conv = upload(client)
    cols = conv["table_columns"]
    for c in cols:
        c["include"] = c["key"] in {"date", "montant"}
    client.put(f"/api/conversions/{conv['id']}/columns", json={"columns": cols})
    text = client.get(f"/api/conversions/{conv['id']}/export?format=csv-fr").content.decode("utf-8-sig")
    assert text.splitlines()[0] == "Date;Montant"


def test_tab_status_and_send(client: TestClient, settings: Settings) -> None:
    status = client.get("/api/tab").json()
    assert status["found"] is False and "introuvable" in status["detail"]
    root = make_tab(settings, projects=2)
    status = client.get("/api/tab").json()
    assert status["found"] is True
    assert [(p["id"], p["name"], p["workspace"]) for p in status["projects"]] == [
        ("proj_0", "Budget 0", "Maison"),
        ("proj_1", "Budget 1", "Maison"),
    ]
    conv = upload(client)
    r = client.post(f"/api/conversions/{conv['id']}/send-to-tab", json={"project_id": "proj_1", "table_name": "Mars"})
    assert r.status_code == 200, r.text
    result = r.json()
    assert result["project_name"] == "Budget 1" and result["rows"] == 18 and result["columns"] == 6
    written = json.loads((root / result["file"]).read_text(encoding="utf-8"))
    assert written["projectId"] == "proj_1" and written["name"] == "Mars" and written["id"] == result["table_id"]
    fields = [written["fields"][f] for f in written["fieldOrder"]]
    assert [(f["label"], f["type"]) for f in fields] == [
        ("Date", "date"),
        ("Date de valeur", "date"),
        ("Libellé", "text"),
        ("Débit", "currency"),
        ("Crédit", "currency"),
        ("Montant", "currency"),
    ]
    assert fields[2]["atelierMapping"] == "entity_title" and fields[3]["options"] == {"currency": "EUR", "decimals": 2}
    first = written["records"][written["rowOrder"][0]]["cells"]
    assert first[written["fieldOrder"][0]] == "2026-03-02" and first[written["fieldOrder"][3]] == 84.2
    assert written["fieldOrder"][4] not in first  # crédit vide : cellule absente
    # atelier.json n'est jamais modifié ; aucun fichier temporaire ne reste.
    assert json.loads((root / "atelier.json").read_text())["projects"]["proj_1"]["tableIds"] == []
    assert [p.name for p in (root / "tables").iterdir()] == [f"{result['table_id']}.json"]
    sent = client.get(f"/api/conversions/{conv['id']}").json()["sent"]
    assert sent[0]["project_name"] == "Budget 1" and sent[0]["table_name"] == "Mars"

    assert client.post(f"/api/conversions/{conv['id']}/send-to-tab", json={"project_id": "proj_9"}).status_code == 404


def test_send_to_tab_validates_with_tab_schema(client: TestClient, settings: Settings) -> None:
    """La table écrite respecte le schéma de fichier de Tab (src/storage/format.ts) : forme et types des cellules."""
    root = make_tab(settings)
    conv = upload(client, PDF, "releve.pdf")
    result = client.post(f"/api/conversions/{conv['id']}/send-to-tab", json={"project_id": "proj_0"}).json()
    table = json.loads((root / result["file"]).read_text(encoding="utf-8"))
    assert set(table) == {
        "id",
        "projectId",
        "name",
        "fieldOrder",
        "rowOrder",
        "fields",
        "records",
        "views",
        "activeViewId",
    }
    assert table["id"].startswith("tbl_") and len(table["id"]) == 14
    allowed_types = {"text", "number", "currency", "percent", "select", "date", "checkbox", "owner", "formula"}
    for fid, field in table["fields"].items():
        assert field["id"] == fid and field["type"] in allowed_types
    for rid in table["rowOrder"]:
        record = table["records"][rid]
        assert record["kind"] == "data"
        assert all(isinstance(v, str | int | float | bool) for v in record["cells"].values())
    assert table["views"][0]["mode"] == "hybrid" and table["activeViewId"] == table["views"][0]["id"]
    assert table["name"] == "releve"


class FakeAI:
    """AIService simulé pour « Structurer avec l'IA »."""

    def __init__(self, online: bool = True, reply: dict[str, Any] | None = None) -> None:
        self.online = online
        self.reply = reply or {
            "columns": [
                {"label": "Date", "type": "date"},
                {"label": "Libellé", "type": "text"},
                {"label": "Débit", "type": "currency"},
            ],
            "rows": [["2026-03-02", "EDF", 84.2], ["2026-03-03", "Boulangerie", 7.9]],
        }
        self.messages: list[ChatMessage] = []
        self.system: str | None = None

    async def chat_json(self, choice: str, messages: list[ChatMessage], system: str | None, schema: dict[str, Any]):
        if choice == "anthropic" and not self.online:
            raise AIProviderError("anthropic", "Mode hors ligne", 403)
        self.messages, self.system = messages, system

        class P:
            name: ProviderName = "anthropic" if choice == "anthropic" else "ollama"
            model = "test"

        return P(), self.reply


def test_structure_with_ai(client: TestClient, override_ai) -> None:  # type: ignore[no-untyped-def]
    fake = FakeAI()
    override_ai(fake)
    conv = upload(client, PDF, "releve.pdf", pdf_mode="text")
    r = client.post(f"/api/conversions/{conv['id']}/structure", json={"provider": "ollama", "instructions": "Relevé"})
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["provider"] == "ollama" and body["sent_lines"] > 10
    view = body["conversion"]
    assert view["detected"] == "Structuré par l'IA (Ollama) · 2 ligne(s)"
    assert [c["key"] for c in view["table_columns"]] == ["date", "libelle", "debit", "credit", "montant"]
    assert view["preview"][0]["montant"] == -84.2
    # Le texte voyage dans la conversation, délimité ; le prompt système ne contient que des consignes.
    assert "<donnees>" in fake.messages[0].content and "RELEVÉ DE COMPTE" in fake.messages[0].content
    assert fake.system is not None and "RELEVÉ" not in fake.system
    assert view["has_text"] is True  # le texte reste disponible pour un nouvel essai


def test_structure_refused_offline_for_claude(client: TestClient, override_ai) -> None:  # type: ignore[no-untyped-def]
    override_ai(FakeAI(online=False))
    conv = upload(client, PDF, "releve.pdf", pdf_mode="text")
    r = client.post(f"/api/conversions/{conv['id']}/structure", json={"provider": "anthropic"})
    assert r.status_code == 403 and r.json()["provider"] == "anthropic"


def test_structure_without_text(client: TestClient, override_ai) -> None:  # type: ignore[no-untyped-def]
    override_ai(FakeAI())
    conv = upload(client)
    assert client.post(f"/api/conversions/{conv['id']}/structure", json={}).status_code == 422


def test_structure_bad_ai_reply(client: TestClient, override_ai) -> None:  # type: ignore[no-untyped-def]
    override_ai(FakeAI(reply={"columns": [], "rows": []}))
    conv = upload(client, PDF, "releve.pdf", pdf_mode="text")
    assert client.post(f"/api/conversions/{conv['id']}/structure", json={}).status_code == 502


def test_samples(client: TestClient) -> None:
    names = [s["name"] for s in client.get("/api/samples").json()]
    assert names == ["releve-mars-2026.pdf", "export-banque.csv", "operations.ofx", "releve-scanne.pdf"]
    r = client.post("/api/samples/releve-mars-2026.pdf")
    assert r.status_code == 201 and r.json()["statement"]["balance_check"] == "ok"
    assert client.post("/api/samples/inconnu.csv").status_code == 404


def test_settings_and_info(client: TestClient, settings: Settings) -> None:
    assert client.get("/api/settings").json()["online_mode"] is False
    assert client.put("/api/settings", json={"online_mode": True}).json()["online_mode"] is True
    info = client.get("/api/settings/info").json()
    assert info["ai_max_lines"] == 200 and ".pdf" in info["extensions"]
    assert info["tab_directory"] == str(settings.tab_dir)


@pytest.mark.parametrize("path", ["/api/health", "/api/ai/providers"])
def test_basic_routes(client: TestClient, path: str) -> None:
    assert client.get(path).status_code == 200
