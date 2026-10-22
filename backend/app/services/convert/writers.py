"""Sorties d'une conversion : CSV (français ou standard), TXT tabulé, XLSX, `atelier-table/1`.

Seules les colonnes incluses sont écrites, dans l'ordre de la table.
"""

from __future__ import annotations

import csv
import io
from datetime import date
from typing import Any

from app.schemas.convert import CellValue, Column, ExportFormat
from app.services.convert.values import format_date_fr

MEDIA_TYPES: dict[ExportFormat, str] = {
    "csv-fr": "text/csv; charset=utf-8",
    "csv": "text/csv; charset=utf-8",
    "txt": "text/plain; charset=utf-8",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "json": "application/json",
}
EXTENSION: dict[ExportFormat, str] = {"csv-fr": "csv", "csv": "csv", "txt": "txt", "xlsx": "xlsx", "json": "json"}

ATELIER_TYPES = {"text": "string", "number": "number", "currency": "number", "date": "date", "boolean": "boolean"}


def _fr_number(value: float, money: bool) -> str:
    text = f"{value:.2f}" if money else _plain(value)
    return text.replace(".", ",")


def _plain(value: float) -> str:
    if float(value).is_integer():
        return str(int(value))
    return f"{value:.10f}".rstrip("0").rstrip(".")


def _fmt(value: CellValue, col: Column, french: bool) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return ("oui" if value else "non") if french else ("true" if value else "false")
    if isinstance(value, int | float) and col.type in {"number", "currency"}:
        return (
            _fr_number(float(value), col.type == "currency")
            if french
            else (f"{float(value):.2f}" if col.type == "currency" else _plain(float(value)))
        )
    if col.type == "date" and isinstance(value, str) and french and len(value) == 10 and value[4] == "-":
        return format_date_fr(value)
    return str(value)


def to_delimited(columns: list[Column], rows: list[dict[str, CellValue]], fmt: ExportFormat) -> bytes:
    """CSV ou TXT. csv-fr : « ; », virgule décimale, JJ/MM/AAAA, UTF-8 avec BOM ; txt : tabulations, à la française."""
    cols = [c for c in columns if c.include]
    french = fmt in {"csv-fr", "txt"}
    delimiter = {"csv-fr": ";", "csv": ",", "txt": "\t"}[fmt]
    buf = io.StringIO()
    writer = csv.writer(buf, delimiter=delimiter, lineterminator="\r\n", quoting=csv.QUOTE_MINIMAL)
    writer.writerow([c.label for c in cols])
    for row in rows:
        writer.writerow([_fmt(row.get(c.key), c, french) for c in cols])
    text = buf.getvalue()
    return (("﻿" if fmt == "csv-fr" else "") + text).encode("utf-8")


def to_xlsx(columns: list[Column], rows: list[dict[str, CellValue]], title: str) -> bytes:
    """Classeur Excel typé : dates en dates, montants au format « 1 234,56 € », en-tête figé et filtrable."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    cols = [c for c in columns if c.include]
    wb = Workbook()
    ws = wb.active
    ws.title = (title or "Conversion")[:31].translate({ord(c): " " for c in "[]:*?/\\"}) or "Conversion"
    ws.append([c.label for c in cols])
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="F3EBD6")
        cell.alignment = Alignment(vertical="center")
    for row in rows:
        out: list[Any] = []
        for c in cols:
            v = row.get(c.key)
            if c.type == "date" and isinstance(v, str) and len(v) == 10:
                try:
                    v = date.fromisoformat(v)
                except ValueError:
                    pass
            out.append(v)
        ws.append(out)
    for i, c in enumerate(cols, start=1):
        letter = get_column_letter(i)
        fmt = {"currency": '#,##0.00 "€";-#,##0.00 "€"', "date": "DD/MM/YYYY", "number": "General"}.get(c.type)
        if fmt:
            for cell in ws[letter][1:]:
                cell.number_format = fmt
        width = max([len(c.label)] + [len(str(r.get(c.key) or "")) for r in rows[:500]])
        ws.column_dimensions[letter].width = min(max(10, width + 2), 60)
    ws.freeze_panes = "A2"
    if cols:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{max(1, len(rows) + 1)}"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def to_atelier_table(name: str, label: str, columns: list[Column], rows: list[dict[str, CellValue]]) -> dict[str, Any]:
    """Table d'échange `atelier-table/1` (docs/SUITE.md §7.2)."""
    cols = [c for c in columns if c.include]
    return {
        "format": "atelier-table/1",
        "name": name,
        "label": label,
        "columns": [{"key": c.key, "label": c.label, "type": ATELIER_TYPES[c.type]} for c in cols],
        "rows": [{c.key: _json_value(row.get(c.key), c) for c in cols} for row in rows],
        "total": len(rows),
        "truncated": False,
    }


def _json_value(value: CellValue, col: Column) -> CellValue:
    if col.type in {"number", "currency"} and not isinstance(value, int | float):
        return None
    if col.type == "boolean" and not isinstance(value, bool):
        return None
    return value
